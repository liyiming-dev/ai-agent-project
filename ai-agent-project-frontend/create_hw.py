import sys
import os
import requests

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "requests"])
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn

def get_desktop_path():
    import ctypes
    from ctypes import wintypes
    CSIDL_DESKTOP = 0x0000
    path_buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
    ctypes.windll.shell32.SHGetFolderPathW(0, CSIDL_DESKTOP, 0, 0, path_buf)
    return path_buf.value

def get_diagram_image():
    plantuml_source = """
@startuml
left to right direction
skinparam packageStyle rectangle

actor "买家" as buyer
actor "卖家" as seller
actor "客服小二" as admin
actor "支付宝业务系统" as pay

package "闲鱼APP系统" {
  usecase "搜索与浏览物品" as UC2
  usecase "沟通与议价" as UC3
  usecase "购买闲置物品" as UC4
  usecase "确认收货" as UC5
  usecase "评价交易" as UC6
  usecase "申请退款与纠纷处理" as UC7
  usecase "发布闲置物品" as UC1
  usecase "管理已发布物品" as UC8
}

buyer --> UC2
buyer --> UC3
buyer --> UC4
buyer --> UC5
buyer --> UC6
buyer --> UC7

seller --> UC1
seller --> UC8
seller --> UC3
seller --> UC6
seller --> UC7

UC5 --> pay : 结算打款
UC4 --> pay : 冻结预付款

admin --> UC7 : 介入处理

@enduml
"""
    payload = {
        "diagram_source": plantuml_source,
        "diagram_type": "plantuml",
        "output_format": "png"
    }
    
    try:
        r = requests.post("https://kroki.io", json=payload, timeout=20)
        if r.status_code == 200:
            with open("usecase_diagram.png", "wb") as f:
                f.write(r.content)
            return "usecase_diagram.png"
    except Exception as e:
        print(f"Error fetching diagram: {e}")
    return None

def create_doc():
    doc = Document()
    # 统一使用宋体
    doc.styles['Normal'].font.name = u'宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    
    heading = doc.add_heading('软件需求第六周作业', 0)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    doc.add_paragraph('项目名称：', style='List Bullet').add_run('闲鱼APP').bold = True
    
    doc.add_heading('1. 业务规则列表 (Business Rules)', level=1)
    doc.add_paragraph('以下是与闲鱼APP项目相关的主要业务规则：')
    
    # 业务规则表
    br_table = doc.add_table(rows=1, cols=3)
    br_table.style = 'Table Grid'
    hdr_cells = br_table.rows[0].cells
    hdr_cells[0].text = '规则编号'
    hdr_cells[1].text = '规则名称'
    hdr_cells[2].text = '规则描述'
    
    rules = [
        ('BR-01', '实名认证限制', '用户必须完成支付宝实名认证及活体认证才能发布商品或下订单。'),
        ('BR-02', '图片数量限制', '发布闲置物品时，单个商品最多允许上传9张图片。'),
        ('BR-03', '资金担保机制', '买家付款后资金由支付宝担保，买家确认收货后（或15天系统自动确认收货后），资金才会结算打入卖家账户。'),
        ('BR-04', '支付超时自动取消', '买家拍下商品生成订单后，如果在30分钟内未完成付款，系统将自动关闭该订单，商品恢复为可售状态。'),
        ('BR-05', '站外引流限制', '用户间的沟通内容中不可包含外部通讯软件（如微信、QQ等）的联系方式及其他交易平台的链接，由于防诈骗机制会触发屏蔽或限制账号。'),
        ('BR-06', '发布数量限制', '为避免商家批量铺货或恶意广告，普通用户每日发布的商品数量上限受到限制（例如最高50个/天）。'),
        ('BR-07', '评价时效规定', '交易成功（确认收货）后，买卖双方有15天的时间进行互评，逾期未评价则系统将对无评价一方自动给出默认好评。')
    ]
    for br_id, br_name, br_desc in rules:
        row_cells = br_table.add_row().cells
        row_cells[0].text = br_id
        row_cells[1].text = br_name
        row_cells[2].text = br_desc
        
    doc.add_paragraph('')
    
    doc.add_heading('2. 用例列表及业务规则引用表', level=1)
    doc.add_paragraph('下表展示了闲鱼系统的核心用例列表，并在最后一列将上述业务规则引用回对应的用例之中：')
    
    # 用例表
    uc_table = doc.add_table(rows=1, cols=5)
    uc_table.style = 'Table Grid'
    hdr_cells = uc_table.rows[0].cells
    hdr_cells[0].text = '用例编号'
    hdr_cells[1].text = '用例名称'
    hdr_cells[2].text = '描述'
    hdr_cells[3].text = '参与者'
    hdr_cells[4].text = '引用的业务规则'
    
    usecases = [
        ('UC-01', '发布闲置物品', '卖家填写闲置物品的信息、价格及运费方式并上传图片进行发布。', '卖家', 'BR-01, BR-02, BR-06'),
        ('UC-02', '搜索与浏览物品', '买家通过关键词搜索或在首页信息流中浏览感兴趣的闲置物品。', '买家', '无'),
        ('UC-03', '沟通与议价', '买卖双方通过内置聊天系统对商品细节进行沟通或发起议价。', '买家, 卖家', 'BR-05'),
        ('UC-04', '购买闲置物品', '买家确定购买后点击“我想要”或立即购买，并完成担保付款。', '买家, 支付宝系统', 'BR-01, BR-04'),
        ('UC-05', '确认收货', '买家收到实物并检查无误后，在订单确认收货，触发担保资金结算给卖家。', '买家, 支付宝系统', 'BR-03'),
        ('UC-06', '评价交易', '交易完成后，买卖两方对此次交易情况和用户体验给出评价。', '买家, 卖家', 'BR-07'),
        ('UC-07', '处理退款与纠纷', '买家申请退款，卖家处理同意或拒绝，若多次协商失败则交由闲鱼客服小二介入。', '买家, 卖家, 客服小二', 'BR-03')
    ]
    for u_id, u_name, u_desc, u_actor, u_rules in usecases:
        row_cells = uc_table.add_row().cells
        row_cells[0].text = u_id
        row_cells[1].text = u_name
        row_cells[2].text = u_desc
        row_cells[3].text = u_actor
        row_cells[4].text = u_rules
        
    doc.add_paragraph('')
    
    doc.add_heading('3. 系统用例图', level=1)
    doc.add_paragraph('下图为闲鱼APP系统基于PlantUML绘制的核心功能用例图：')
    
    img_path = get_diagram_image()
    if img_path and os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6.0))
        p = doc.add_paragraph('图1：闲鱼APP系统用例图', style='Caption')
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        os.remove(img_path)
    else:
        doc.add_paragraph('（图片在线生成失败，由于网络等原因未成功获取，可以用PlantUML插件等重新导入）')
        
    desktop = get_desktop_path()
    save_path = os.path.join(desktop, '软件需求第六周作业.docx')
    doc.save(save_path)
    print(save_path)

if __name__ == '__main__':
    create_doc()
